#!/usr/bin/env python3
"""Generate original deterministic DOCX/XLSX/PPTX fixtures with pinned tools.

Use python-docx==1.2.0, openpyxl==3.1.5, python-pptx==1.0.2.
The fixture-tool lock records transitive versions; no tool is a runtime dependency.
"""
import argparse
from datetime import datetime
from io import BytesIO
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

STAMP = datetime(2020, 1, 1)


def stable_archive(raw, target):
    """Normalize ZIP metadata and generated core timestamps for repeatability."""
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(raw)) as source, zipfile.ZipFile(output, 'w', compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for name in sorted(source.namelist()):
            data = source.read(name)
            if name == 'docProps/core.xml':
                core = ET.fromstring(data)
                for element in core:
                    if element.tag.rsplit('}', 1)[-1] in ('created', 'modified'):
                        element.text = '2020-01-01T00:00:00Z'
                data = ET.tostring(core, encoding='utf-8', xml_declaration=True)
            info = zipfile.ZipInfo(name, (2020, 1, 1, 0, 0, 0))
            info.create_system = 3
            info.external_attr = 0o600 << 16
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, data, compresslevel=9)
    target.write_bytes(output.getvalue())


def documents(output):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    document = Document()
    document.core_properties.title = 'Repository Handbook'
    document.core_properties.author = 'Fixture Author'
    document.core_properties.created = STAMP
    document.core_properties.modified = STAMP
    document.add_heading('Repository Handbook', 1)
    document.add_paragraph('Index source files and documents into one searchable catalogue.')
    document.add_heading('Recovery', 2)
    document.add_paragraph('After an interrupted update, restart indexing to finish pending work.')
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = 'Task'
    table.rows[0].cells[1].text = 'State'
    for name, state in [('Index', 'Complete'), ('Search', 'Ready')]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = name, state
    raw = BytesIO()
    document.save(raw)
    stable_archive(raw.getvalue(), output / 'handbook.docx')

    workbook = Workbook()
    workbook.properties.title = 'Repository Inventory'
    workbook.properties.creator = 'Fixture Author'
    workbook.properties.created = STAMP
    workbook.properties.modified = STAMP
    sheet = workbook.active
    sheet.title = 'Documents'
    for row in [('Path', 'Kind', 'Count'), ('manual.pdf', 'PDF', 2), ('guide.md', 'Markdown', 3)]:
        sheet.append(row)
    sheet = workbook.create_sheet('Operations')
    for row in [('Command', 'Purpose'), ('index', 'Refresh changed sources'), ('search', 'Find relevant sections')]:
        sheet.append(row)
    raw = BytesIO()
    workbook.save(raw)
    stable_archive(raw.getvalue(), output / 'inventory.xlsx')

    slides = Presentation()
    slides.core_properties.title = 'Search Workflow'
    slides.core_properties.author = 'Fixture Author'
    slides.core_properties.created = STAMP
    slides.core_properties.modified = STAMP
    for heading, lines in [('Indexing', ['Read code and documents.', 'Keep page and heading context.']), ('Searching', ['Combine relevant evidence.', 'Open the original source.'])]:
        slide = slides.slides.add_slide(slides.slide_layouts[1])
        slide.shapes.title.text = heading
        text = slide.placeholders[1].text_frame
        text.paragraphs[0].text = lines[0]
        text.add_paragraph().text = lines[1]
    raw = BytesIO()
    slides.save(raw)
    stable_archive(raw.getvalue(), output / 'workflow.pptx')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=Path('fixtures/docs'))
    parser.add_argument('--encrypted', action='store_true', help='also regenerate the randomized encrypted DOCX (msoffcrypto-tool==5.4.2)')
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    documents(args.output)
    if args.encrypted:
        import msoffcrypto
        plaintext = (args.output / 'handbook.docx').read_bytes()
        encrypted = BytesIO()
        msoffcrypto.OfficeFile(BytesIO(plaintext)).encrypt('fixture-only', encrypted)
        encoded = encrypted.getvalue()
        reader = msoffcrypto.OfficeFile(BytesIO(encoded))
        reader.load_key(password='fixture-only', verify_password=True)
        decoded = BytesIO()
        reader.decrypt(decoded, verify_integrity=True)
        if decoded.getvalue() != plaintext:
            raise RuntimeError('encrypted fixture round trip failed')
        (args.output / 'encrypted.docx').write_bytes(encoded)


if __name__ == '__main__':
    main()
