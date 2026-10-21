#!/usr/bin/env python3
"""Generate the 7 DOCX, 7 XLSX and 6 PPTX corpus cases from pinned Markdown excerpts."""
import argparse
from datetime import datetime
import hashlib
from io import BytesIO
import importlib.util
import json
from pathlib import Path
from unittest.mock import patch


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--inputs', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    inputs = json.loads(args.inputs.read_text())
    assert len(inputs) == 20
    spec = importlib.util.spec_from_file_location('office_fixture', Path(__file__).with_name('make-office-fixtures.py'))
    helper = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(helper)
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    import msoffcrypto
    from msoffcrypto.method import ecma376_agile
    from msoffcrypto.method.container import ecma376_encrypted
    args.output.mkdir(parents=True, exist_ok=True)
    for index, source in enumerate(inputs):
        if index < 7:
            extension = 'docx'
            document = Document()
            document.add_heading(source['title'], 0)
            document.add_heading('Source excerpt', 1)
            document.add_paragraph(source['text'])
            document.add_heading('Attribution', 1)
            document.add_paragraph(source['url'])
            buffer = BytesIO(); document.save(buffer)
        elif index < 14:
            extension = 'xlsx'
            workbook = Workbook()
            sheet = workbook.active; sheet.title = 'Source excerpt'
            sheet.append(['Title', source['title']])
            sheet.append(['Source', source['url']])
            for number, sentence in enumerate(source['text'].split('. '), 1):
                sheet.append([number, sentence])
            buffer = BytesIO(); workbook.save(buffer)
        else:
            extension = 'pptx'
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = source['title']
            slide.placeholders[1].text = source['text']
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = 'Attribution'
            slide.placeholders[1].text = source['url']
            buffer = BytesIO(); presentation.save(buffer)
        target = args.output / f'office-{index+1:02d}.{extension}'
        helper.stable_archive(buffer.getvalue(), target)
        if index == 6:
            # Public test data only: fixed entropy and OLE timestamps make this
            # deliberately encrypted rejection fixture reproducible, not secret.
            counter = 0
            def fixture_bytes(size):
                nonlocal counter
                data = bytearray()
                while len(data) < size:
                    data.extend(hashlib.sha256(b'dkc public encrypted fixture v1' + counter.to_bytes(8, 'little')).digest())
                    counter += 1
                return bytes(data[:size])
            class FixedDateTime(datetime):
                @classmethod
                def now(cls, tz=None):
                    return cls(2020, 1, 1, tzinfo=tz)
            plaintext = target.read_bytes()
            encrypted = BytesIO()
            with patch.object(ecma376_agile, '_random_buffer', fixture_bytes), patch.object(ecma376_encrypted, 'datetime', FixedDateTime):
                msoffcrypto.OfficeFile(BytesIO(plaintext)).encrypt('fixture-only', encrypted)
            reader = msoffcrypto.OfficeFile(BytesIO(encrypted.getvalue()))
            reader.load_key(password='fixture-only', verify_password=True)
            decoded = BytesIO(); reader.decrypt(decoded, verify_integrity=True)
            assert decoded.getvalue() == plaintext
            target.write_bytes(encrypted.getvalue())
        if index == 13:
            # Preserve the ZIP prefix but remove its end/central directory.
            target.write_bytes(target.read_bytes()[:128])
        print(target.name, hashlib.sha256(target.read_bytes()).hexdigest(), flush=True)


if __name__ == '__main__':
    main()
